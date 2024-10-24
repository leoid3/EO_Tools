import csv
from config import *
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from functions.calcul import simulation_time
from functions.find_tm import centroid
from functions.revisit import revisit_over_a_latitude, revi

def save_gs_visibility(result, name):
    """
    Permet de sauvegarder les resultats des GS dans un fichier csv.
    """
    if len(result) == 0:
        print("NO OPPORTUNITIES")
    else:
        filename = result_folder / name / f"Result Ground Station visibility ({result[0][0]}).csv"
        with open(filename, 'w') as csvfile:
            csvwriter = csv.writer(csvfile)
            csvwriter.writerow(resultfields_gs)
            csvwriter.writerows(result)

def save_poi_visibility(result, name):
    """
    Permet de sauvegarder les resultats des POIs dans un fichier csv.
    """
    if len(result) == 0:
        print("NO OPPORTUNITIES")
    else:
        filename = result_folder / name / f"Result POI visibility ({result[0][0]}).csv"
        with open(filename, 'w') as csvfile:
            csvwriter = csv.writer(csvfile)
            csvwriter.writerow(resultfields_poi)
            csvwriter.writerows(result)

def general_result(mission, gs_opportunities, poi_opportunities):
    """
    Permet de creer le mission report.pdf.
    """
    # Function to replace markup with actual values
    def replace_placeholders(text, values):
        """
        Permet de remplacer les placeholder par les données de la simulation.
        """
        for key, value in values.items():
            if value== poi:
                new_value =''
                for i in range(len(value)):
                    if i==0:
                        temp = value[i]
                        new_value ="- " + temp.get_name()
                    else:
                        temp = value[i]
                        new_value =new_value + "<br></br>" + "- "+ temp.get_name()
                text = text.replace(key, new_value)
            elif value== gs:
                new_value =''
                for i in range(len(value)):
                    if i==0:
                        temp = value[i]
                        new_value ="- " + temp.get_name()
                    else:
                        temp = value[i]
                        new_value =new_value + "<br></br>" + "- "+ temp.get_name()
                text = text.replace(key, new_value)
            elif value== listpoi:
                new_value =''
                for i in range(len(value)):
                    temp = value[i]
                    if i==0:
                        pass
                    else:
                        if temp.IsArea() == True:
                            long, lat = centroid(temp.get_area())
                        else:
                            lat = temp.get_coordinate(0)[0]
                            long = temp.get_coordinate(0)[1]
                        if poi_visi[i-1] == False:
                            number = "Not available for an area"
                        else:
                            number = poi_visi[i-1]
                        if i == 1:
                            new_value = temp.get_name() + "<br></br>" + "Coordinates : " + str(lat) + "°  " + str(long) + "°" + "<br></br>" + "Altitude : " + str(temp.get_altitude()) + " m" + "<br></br>" + "Total number of times visible : " + str(number) + "<br></br>" + "Mean duration time : " + "XX" +" s" + "<br></br>"+ "<POIopportunities>" + "<br></br>"
                        else:
                            new_value = new_value + "<br></br>" + "Name of the point of interest : " + temp.get_name() + "<br></br>" + "Coordinates : " + str(lat) + "°  " + str(long) + "°" + "<br></br>" + "Altitude : " + str(temp.get_altitude()) + " m" + "<br></br>" + "Total number of times visible : " + str(number) + "<br></br>" + "Mean duration time : " + "XX" +" s" + "<br></br>"+ "<POIopportunities>" + "<br></br>"
                text = text.replace(key, new_value)
            elif value == listegs:
                new_value =''
                for i in range(len(value)):
                    temp = value[i]
                    if i== 0:
                        pass
                    else:
                        lat, long = temp.get_coordinate()
                        number = gs_visi[i-1]
                        if i == 1:
                            new_value = temp.get_name() + "<br></br>" + "Coordinates : " + str(lat) + "°  " + str(long) + "°" + "<br></br>" + "Altitude : " + str(temp.get_altitude()) + " m" + "<br></br>" + "Antenna characteristics :" + "<br></br>" + "- Elevation : " + str(temp.get_elevation()) + " °" + "<br></br>" + "- Band : " + temp.get_band() + "<br></br>" + "- Debit : " + str(temp.get_debit()) + " Mbps" + "<br></br>" + "Total number of opportunities : " + str(number) + "<br></br>" + "Mean duration time : " + "XX" + " s" + "<br></br>" + "Mean data size send/received : " + "XX" + " Mb" + "<br></br>" + "<GSopportunities>" + "<br></br>"
                        else:
                            new_value = new_value + "<br></br>" + "Name of the ground station : " + temp.get_name() + "<br></br>" + "Coordinates : " + str(lat) + "°  " + str(long) + "°" + "<br></br>" + "Altitude : " + str(temp.get_altitude()) + " m" + "<br></br>" + "Antenna characteristics :" + "<br></br>" + "- Elevation : " + str(temp.get_elevation()) + " °" + "<br></br>" + "- Band : " + temp.get_band() + "<br></br>" + "- Debit : " + str(temp.get_debit()) + " Mbps" + "<br></br>" + "Total number of opportunities : " + str(number) + "<br></br>" + "Mean duration time : " + "XX" + " s" + "<br></br>" + "Mean data size send/received : " + "XX" + " Mb" + "<br></br>" + "<GSopportunities>" + "<br></br>"
                text = text.replace(key, new_value)
            else:
                text = text.replace(key, value)
        return text

    def plot_result_gs(opportunities):
        """
        Permet de creer le graph des GS.
        """
        fig2d = plt.figure(figsize=(10, 6))
        ax_2D = fig2d.add_subplot(111)
        temp = opportunities
        abscisse = []
        ordonnees = []
        for i in range(len(temp)):
            if i== 0:
                title = temp[0]
            else:
                for j in range(len(temp[i])):
                    data = temp[i][j]
                    abscisse.append(data[0])
                    ordonnees.append(data[1])
        ax_2D.bar(abscisse, ordonnees, width=0.35, color=list_colors, align='center')
        ax_2D.set_title("Number of visibility per satellite at " + title)
        ax_2D.legend()
        fig2d.savefig(result_folder / mission.get_name() / f"Visibility at {title}.png", bbox_inches='tight')
        img = Image(result_folder / mission.get_name() / f"Visibility at {title}.png", width=400, height=300)
        return img

    def plot_result_poi(opportunities):
        """
        Permet de creer le graph des POIs.
        """
        temp = opportunities
        
        if temp[1][0][0]:
            visi_list = []
            color = []
            fig, ax = plt.subplots(figsize=(15, 15))
            if temp[1][1][4]:
                grid_poi = temp[1][1][3]
                poi_poly = temp[1][1][2]
                for k in range(len(grid_poi)):
                    if grid_poi[k] == None:
                        pass
                    else:
                        gpd.GeoSeries(grid_poi[k]).boundary.plot(ax=ax)
                        gpd.GeoSeries([poi_poly[k]]).boundary.plot(ax=ax,color="red")
            else:
                grid_poi = temp[1][1][3]
                poi_poly = temp[1][1][2]
                gpd.GeoSeries(grid_poi).boundary.plot(ax=ax)
                gpd.GeoSeries([poi_poly[0]]).boundary.plot(ax=ax,color="red")
            for i in range(len(temp)):
                for j in range(len(temp[i])):
                    if i== 0:
                        title = temp[0]
                    else:
                        visi_list.append(temp[i][j][1])
            for i in range(len(visi_list)):
                color_temp = []
                for j in range(len(visi_list[i])):
                        if i ==0:
                            if visi_list[0][j][0]:
                                if visi_list[0][j][1]:
                                    color_temp.append([2, visi_list[0][j][2]])
                                else:
                                    color_temp.append([1, visi_list[0][j][2]])
                            else:
                                color_temp.append([0, visi_list[0][j][2]])
                        else:
                            print(color[0][j][0])
                            if color[0][j][0]==2:
                                pass
                            elif color[0][j][0]==1:
                                if visi_list[i][j][0]:
                                    if visi_list[i][j][1]:
                                        color[0][j][0]=2
                                    else:
                                        pass
                                else:
                                    pass
                            elif color[0][j][0]==0:
                                print(visi_list[i][j][0])
                                if visi_list[i][j][0]:
                                    if visi_list[i][j][1]:
                                        color[0][j][0]=2
                                    else:
                                        color[0][j][0]=1
                                else:
                                    pass
                if i==0:
                    color.append(color_temp)
            prct = 0
            for i in range(len(color)):
                for j in range(len(color[i])):
                    if color[i][j][0]==2:
                        gpd.GeoSeries(color[i][j][1]).plot(ax=ax, color="green")
                        prct += 1
                    elif color[i][j][0]==1:
                        gpd.GeoSeries(color[i][j][1]).plot(ax=ax, color="orange")
                    else:
                        gpd.GeoSeries(color[i][j][1]).plot(ax=ax, color="red")
            green = Line2D([0], [0], marker='o', color='w', markerfacecolor='green', markersize=10, label='Observation possible')
            orange = Line2D([0], [0], marker='o', color='w', markerfacecolor='orange', markersize=10, label='Area visible but minimum sza not reached')
            red =Line2D([0], [0], marker='o', color='w', markerfacecolor='red', markersize=10, label='Observation impossible')
            prct = (prct*100)/len(color[0])
            ax.legend(handles=[green, orange, red])
            ax.set_title(f"{prct:.2f}% of {title} covered during the simulation ")
            fig.savefig(result_folder / mission.get_name() / f"Coverage map of {title}.png", bbox_inches='tight')
            img = Image(result_folder / mission.get_name() / f"coverage map of {title}.png", width=400, height=300)
            return img
        else:
            fig2d = plt.figure(figsize=(10, 6))
            ax_2D = fig2d.add_subplot(111)
            abscisse = []
            ordonnees = []
            for i in range(len(temp)):
                if i== 0:
                    title = temp[0]
                else:
                    for j in range(len(temp[i])):
                        data = temp[i][j]
                        abscisse.append(data[1])
                        ordonnees.append(data[2])
            ax_2D.bar(abscisse, ordonnees, width=0.35, color=list_colors, align='center')
            ax_2D.set_title("Number of visibility per satellite at " + title)
            ax_2D.legend()
            fig2d.savefig(result_folder / mission.get_name() / f"Visibility at {title}.png", bbox_inches='tight')
            img = Image(result_folder / mission.get_name() / f"Visibility at {title}.png", width=400, height=300)
            return img
    
    def number_of_visibility(opportunities):
        temp = opportunities
        total = 0
        for i in range(len(temp)):
            if i== 0:
                pass
            else:
                for j in range(len(temp[i])):
                    data = temp[i][j]
                    total = total + data[1]
        return total
    ##### Mission
    missionname = mission.get_name()
    missiontype = mission.get_type()
    delta, _, _, _ = simulation_time(mission)
    if delta.days == 0:
        missionduration = 1
    else:
        missionduration = delta.days
    missionsza = mission.get_minsza()
    poi = []
    for i in range(mission.get_nb_poi()):
        poi.append(mission.get_poi(i))
    gs = []
    for i in range(mission.get_nb_gs()):
        gs.append(mission.get_gs(i))
    missionconst = mission.get_constellation()

    ##### Constellation
    constname = missionconst.get_name()
    numbersat = missionconst.get_walkerT()
    numberplane = missionconst.get_walkerP()
    phasing = missionconst.get_walkerF()
    satmodel = missionconst.get_model()
    satmodel_name = satmodel.get_name()
    orbit = satmodel.get_orbit()
    a, e, incl, _, _, _ = orbit.get_all()
    if ((a-earth_radius)/1000>0) and ((a-earth_radius)/1000<=2000):
        orbittype = "LEO"
    elif  ((a-earth_radius)/1000>2000) and ((a-earth_radius)/1000<35786):
        orbittype = "MEO"
    elif (a-earth_radius)/1000==35786:
        orbittype = "GEO"
    else:
        orbittype = "Undefined"
    semimajoraxis = a/1000
    inclination = incl
    eccentricity = e
    period = 2*np.pi*np.sqrt((a**3)/mu)
    numberorbit = 86400/period
    swath = satmodel.get_swath()
    depointing = satmodel.get_depoiting()
    revisit0 = revisit_over_a_latitude(a, swath, depointing, 0, numbersat, incl)
    revisit45=revisit_over_a_latitude(a, swath, depointing, 45, numbersat, incl)
    revisit60=revisit_over_a_latitude(a, swath, depointing, 60, numbersat, incl)
    #revisit0 = revi(a, incl, 0)
    #revisit45 = revi(a, incl, 45)
    #revisit60 = revi(a, incl, 60)

    ##### POI
    listpoi = []
    poi_visi = []
    listpoi.append("liste")
    for i in range(mission.get_nb_poi()):
        listpoi.append(mission.get_poi(i))
    for i in range(len(poi_opportunities)):
        temp = poi_opportunities[i]
        if temp[0]:
            poi_visi.append(False)
        else:
            data = number_of_visibility(temp)
            poi_visi.append(data)

    ##### GS
    listegs = []
    gs_visi = []
    listegs.append("liste")
    for i in range(mission.get_nb_gs()):
        listegs.append(mission.get_gs(i))
    for i in range(len(gs_opportunities)):
        data = number_of_visibility(gs_opportunities[i])
        gs_visi.append(data)


    ##### Open the .docx report template
    doc = Document("Mission report Template.docx")
    text_elements = []
    for para in doc.paragraphs:
        text_elements.append(para.text)

    plot_images = {
        "Map with GS & POI" : result_folder / mission.get_name() / f"Ground Track of {missionname}.png",
        "3D orbit" : result_folder / mission.get_name() / f"Orbit of {missionname}.png",
        "Ground track" : result_folder / mission.get_name() / f"Ground Track of {missionname}.png",
        "<GSopportunities>" : result_folder / f"Ground Track of {missionname}.png",
        "<POIopportunities>" : result_folder / f"Ground Track of {missionname}.png"
    }

    # Define specific keywords that indicate titles and sections
    main_title = ["Mission report"]
    secondary_title = [
        "Mission details",
        "Constellation details",
        "Point of interest details",
        "Ground Station details"
    ]

    # Define the values to replace
    values = {
        
        "<MissionName>": missionname,
        "<MissionType>": missiontype,
        "<MissionDuration>": str(missionduration),
        "<MissionSZA>": str(missionsza),
        "<POIName> (<Latitude> <Longitude>)": poi,
        "<GSName> (<Latitude> <Longitude>)": gs,
        "<MissionConstellation>": constname,
        "<ConstellationName>": constname,
        "<NumberofSat>": str(numbersat), 
        "<NumberofPlane>": str(numberplane), 
        "<PhasingFactor>": str(phasing), 
        "<SatName>": satmodel_name, 
        "<Typeoforbit>": orbittype, 
        "<semimajoraxis>": str(semimajoraxis), 
        "<inclination>": str(inclination), 
        "<eccentricity>": str(eccentricity), 
        "<period>": str(period), 
        "<Numberoforbit>": str(numberorbit), 
        "<revisitat0>": str(revisit0), 
        "<revisitat45>": str(revisit45), 
        "<revisitat60>": str(revisit60), 
        "<swath>": str(swath), 
        "<depointing>": str(depointing),
        "<POIStart>": listpoi,
        "<GSStart>": listegs,   
        "XX" : "PLACEHOLDER",
    }

    # Create a custom style for titles
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='secondarytitle', fontSize=16, leading=16, bold=True))
    styles.add(ParagraphStyle(name='maintitle',fontSize=28, leading=16, alignment=TA_CENTER, spaceAfter=12, bold=True))
    styles.add(ParagraphStyle(name='body',fontSize=12, leading=16))

    # Create the PDF document with enhanced formatting
    pdf_filename = f"Mission Report for {mission.get_name()}.pdf"
    pdf = SimpleDocTemplate(pdf_filename, pagesize=A4)

    elements = []

    for text in text_elements:
        # Replace placeholders
        formatted_text = replace_placeholders(text, values)
        
        # Check if the text matches any of the keywords that indicate a title
        if any(keyword in formatted_text for keyword in main_title):
            elements.append(Paragraph(formatted_text, styles["maintitle"]))   
        elif any(keyword in formatted_text for keyword in secondary_title):
            elements.append(PageBreak())
            elements.append(Paragraph(formatted_text, styles["secondarytitle"]))
        else:
            passed = 0
            
            for keyword, image_path in plot_images.items():
                if keyword in formatted_text:
                    if keyword == "3D orbit":
                        img = Image(image_path, width=400, height=300)
                        elements.append(img)
                        passed = 1
                    elif keyword == "<GSopportunities>":
                        parts = formatted_text.split("<GSopportunities>")
                        for i, part in enumerate(parts):
                            if i<len(parts)-1:
                                elements.append(Paragraph(part, styles["body"]))
                                img = plot_result_gs(gs_opportunities[i])
                                elements.append(img)
                        passed = 1
                        break
                    elif keyword == "<POIopportunities>":
                        parts = formatted_text.split("<POIopportunities>")
                        for i, part in enumerate(parts):
                            if i<len(parts)-1:
                                elements.append(Paragraph(part, styles["body"]))
                                img = plot_result_poi(poi_opportunities[i])
                                elements.append(img)
                        passed = 1
                        break
                    else:
                        img = Image(image_path, width=600, height=300)
                        elements.append(img)
                        passed = 1
            if passed == 0:
                elements.append(Paragraph(formatted_text, styles["body"]))
                    
            
        elements.append(Spacer(1, 12))

    # Build PDF
    pdf.build(elements)

    # Output the path to the generated PDF
    print(f"PDF created successfully: {pdf_filename}") 

  